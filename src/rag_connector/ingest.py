"""The ingest kit: build a corpus you own — **NOT part of the connector contract.**

Everything in this module runs *before* a connector exists. Loading a folder of
documents, splitting them into chunks, and fingerprinting the result is what the
owner of a RAG system does to their own corpus; a connector is a **read-only**
interface onto a system someone else already built that way. Nothing here is
reachable through :class:`~rag_connector.base.RagPipeline`, and nothing here may
become reachable through it — the contract stays read-only (see
``docs/contract.md``, "The pipeline", and
:mod:`rag_connector.reference_provision`).

So this module is declared API for a different audience: a **host building its
own corpus**. The Reference RAG is the first such host, but it is not a special
one — Pelorus and any other product that ingests a folder of documents should
consume these functions rather than write a second chunker. That is the whole
point of promoting them out of ``reference.py``: two independently-written
chunkers over the same 92 documents disagreed on 69 of them, which is a corpus
difference nobody asked for and nobody can see.

Import cost is deliberately near zero. Every optional parser (``pypdf``,
``python-docx``, ``striprtf``) is imported inside the function that needs it, so
``import rag_connector.ingest`` succeeds on a bare core install with no extras —
``tests/test_ingest.py`` asserts exactly that in a clean subprocess.

Stability notes for anyone depending on this:

- ``chunk_id`` is ``{source_identity}:chunk-{doc_chunk_index}`` and
  :func:`document_source_identity` IS the normalized relative path, so an id
  reads ``booking_cancellation_policy.md:chunk-0`` -- self-describing, and the
  same spelling Pelorus already shows on its MCP and connector surfaces. Both
  are load-bearing for every persisted artifact that cites a chunk; changing
  either regrounds every stored citation.
- ``char_start``/``char_end`` are an **invariant, not a hint**:
  ``document.content[char_start:char_end] == chunk.text`` exactly, for every
  chunk. Callers resolve citations back to source with those offsets.
"""

from __future__ import annotations

import importlib.resources
import json
import os
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

from .base import ChunkRecord
from .fingerprints import (
    corpus_fingerprint,
    document_source_fingerprint,
    document_source_identity,
)

__all__ = [
    "BUNDLED_CHUNK_OVERLAP",
    "BUNDLED_CHUNK_SIZE",
    "BUNDLED_CORPORA",
    "CHUNKER_METHOD",
    "CHUNKER_VERSION",
    "DEFAULT_BUNDLED_CORPUS",
    "LoadedDocument",
    "SUPPORTED_SOURCE_EXTENSIONS",
    "build_bundled_chunks",
    "bundled_chunks_path",
    "bundled_corpus_path",
    "chunk_loaded_documents",
    "chunks_from_jsonl",
    "chunks_to_jsonl",
    "corpus_fingerprint",
    "document_source_fingerprint",
    "document_source_identity",
    "load_bundled_chunks",
    "load_bundled_corpus",
    "load_documents",
]

#: Corpora shipped as package data under ``rag_connector/corpus/``. MIT, like
#: the rest of this package -- see that directory's README.
BUNDLED_CORPORA = ("pelorus_space",)

#: The corpus used when a caller asks for "the bundled one" without naming it.
DEFAULT_BUNDLED_CORPUS = "pelorus_space"

#: The chunking parameters the bundled PRE-CHUNKED corpus was produced at, and
#: the only ones at which it can be regenerated. They are constants rather than
#: defaults because the frozen file is a fact about a specific split, not about
#: this kit's preferences -- see :func:`load_bundled_chunks`.
BUNDLED_CHUNK_SIZE = 1000
BUNDLED_CHUNK_OVERLAP = 200

#: The name and version of the splitting behaviour in
#: :func:`chunk_loaded_documents`, recorded in a Dataset manifest's ``chunking``
#: block (see ``docs/dataset-spec.md``). The boundary rule, the trailing-chunk
#: rule and the whitespace rule are behaviour, not parameters, so the same
#: ``chunk_size``/``chunk_overlap`` only reproduce a split under the same
#: version. **Bump the version in the same change that alters any of them**;
#: the frozen-chunk drift test in ``tests/test_ingest.py`` is what notices that
#: the behaviour moved.
CHUNKER_METHOD = "rag-connector-chunker"
CHUNKER_VERSION = "1"

#: File extensions :func:`load_documents` will parse. ``.txt``/``.md`` need no
#: extras; the rest need the ``reference`` extra's parsers.
SUPPORTED_SOURCE_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx", ".rtf"})

#: Separator priority for boundary-aware splitting, longest/strongest first.
_CHUNK_SEPARATORS = ("\n\n", "\n", ". ", "? ", "! ", " ")

#: A chunk may end early to land on a separator, but only inside the last 20% of
#: its window — otherwise a stray newline near the start would halve every chunk.
_BOUNDARY_SEARCH_FRACTION = 0.8


@dataclass
class LoadedDocument:
    filename: str
    filepath: str
    relative_path: str
    file_type: str
    content: str
    source_identity: str
    source_fingerprint: str


@contextmanager
def _parser(extension: str, distribution: str) -> Iterator[None]:
    """Turn a missing optional parser into a message that names the way out.

    A core install has no parsers, and pointing :func:`load_documents` at a
    folder that happens to hold one PDF is an ordinary thing to do. Left
    unguarded that raises ``ModuleNotFoundError: No module named 'pypdf'``,
    naming an import the caller never wrote. Say instead which file type
    needed a parser and which install provides it -- the same shape
    ``reference_provision._open_client`` uses for chromadb.
    """
    try:
        yield
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError(
            f"Reading {extension} files requires the {distribution!r} parser, "
            'which the "reference" extra installs: '
            'pip install "rag-connector[reference]"'
        ) from exc


def _read_file(path: Path) -> str:
    """Parse one source file to text.

    Every parser import is local on purpose: the core install has none of them,
    and a module-level import would make the whole ingest kit unimportable
    without the ``reference`` extra.
    """
    ext = path.suffix.lower()
    # utf-8-sig, not utf-8: a byte-order mark is an encoding artifact, and a
    # plain utf-8 read leaves it in the text as U+FEFF. It would then be the
    # first character of chunk 0, inside the text a citation resolves to and
    # inside every fingerprint over that text -- so one editor that writes a
    # BOM changes the corpus. Files without a BOM read identically.
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8-sig")
    if ext == ".pdf":
        with _parser(".pdf", "pypdf"):
            from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == ".docx":
        with _parser(".docx", "python-docx"):
            from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".rtf":
        with _parser(".rtf", "striprtf"):
            from striprtf.striprtf import rtf_to_text
        return rtf_to_text(path.read_text(encoding="utf-8-sig"))
    raise ValueError(f"Unsupported file type: {ext}")


def _iter_source_files(root: Path) -> Iterator[Path]:
    """Yield every supported source file under ``root``, skipping dot-directories.

    Subfolders yes, dot-subfolders no. A corpus folder is a working directory:
    it holds ``.git``, and for a Pelorus dataset it holds ``.pelorus`` -- state
    the product wrote about the corpus, not documents in it. Ingesting a tool's
    own state back into the corpus it describes is never what the owner asked
    for, and the exclusion is what Pelorus's ``iter_source_files`` already does.

    Walking with :func:`os.walk` rather than ``rglob`` is what makes the skip
    cheap and total: pruning ``dirnames`` in place stops the walk from
    descending at all, at every level, instead of filtering paths after a
    ``.git`` tree has already been enumerated.
    """
    for current, dirnames, filenames in os.walk(root):
        dirnames[:] = [name for name in dirnames if not name.startswith(".")]
        current_path = Path(current)
        for filename in filenames:
            path = current_path / filename
            if path.suffix.lower() in SUPPORTED_SOURCE_EXTENSIONS:
                yield path


def load_documents(folder_path: str) -> list[LoadedDocument]:
    """Load every supported document under ``folder_path``, recursively.

    Dot-prefixed directories (``.git``, ``.pelorus``, and anything else hidden)
    are skipped at every level -- see :func:`_iter_source_files`.

    Results are sorted by path so a corpus fingerprint does not depend on
    filesystem enumeration order -- and sorted by the POSIX *spelling* of the
    relative path, not by the ``Path`` object. Comparing ``Path`` objects sorts
    on a platform-dependent key: Windows casefolds and separates with a
    backslash, POSIX compares the raw string with a forward slash. Both
    ``global_index`` and the corpus fingerprint are derived from this order, so
    the object ordering hands a mixed-case or nested corpus one chunk set on
    Windows and a different one on Linux, from the same documents.
    """
    root = Path(folder_path).resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Folder path does not exist or is not a directory: {folder_path}")
    out: list[LoadedDocument] = []
    ordered = sorted(_iter_source_files(root), key=lambda p: p.relative_to(root).as_posix())
    for path in ordered:
        try:
            content = _read_file(path)
        except UnicodeDecodeError as exc:
            # Still fail-loud, but name the file: the bare decode error aborts
            # the whole folder load without saying which document to fix.
            raise ValueError(
                f"{path} is not valid UTF-8 ({exc}). Re-encode the file as "
                "UTF-8 or remove it from the source folder."
            ) from exc
        file_type = path.suffix.lower().lstrip(".")
        source_identity = document_source_identity(str(root), str(path))
        out.append(LoadedDocument(
            filename=path.name,
            filepath=str(path),
            relative_path=str(path.relative_to(root)).replace("\\", "/"),
            file_type=file_type,
            content=content,
            source_identity=source_identity,
            source_fingerprint=document_source_fingerprint(file_type, content),
        ))
    return out


@contextmanager
def bundled_corpus_path(name: str = DEFAULT_BUNDLED_CORPUS) -> Iterator[str]:
    """Yield a real directory path for one of the corpora shipped in this package.

    A context manager rather than a plain path because that is the honest shape:
    package data is not guaranteed to exist as a file on disk (a zip-imported
    install has to materialise it first), and ``importlib.resources.as_file`` is
    what handles both cases. The path is valid for the duration of the ``with``
    block and no longer.

    Read through :mod:`importlib.resources` for the same reason
    ``CONNECTOR-INFO.md`` is: an editable install resolves to ``src/`` on disk,
    a wheel install does not, and a path relative to this file would work
    locally and fail once installed.
    """
    if name not in BUNDLED_CORPORA:
        raise ValueError(
            f"Unknown bundled corpus {name!r}. Available: "
            f"{', '.join(BUNDLED_CORPORA)}."
        )
    resource = importlib.resources.files(__package__).joinpath("corpus").joinpath(name)
    with importlib.resources.as_file(resource) as path:
        if not Path(path).is_dir():
            # Package data that silently fails to ship is invisible from a
            # green suite -- every local run reads it off src/. Say so here
            # rather than reporting "no supported documents found".
            raise FileNotFoundError(
                f"The bundled {name!r} corpus is not present in this "
                f"installation of rag_connector (looked in {path}). It is "
                "declared in [tool.setuptools.package-data]; a wheel built "
                "without that declaration will not carry it."
            )
        yield str(path)


def load_bundled_corpus(name: str = DEFAULT_BUNDLED_CORPUS) -> list[LoadedDocument]:
    """Load a corpus shipped with this package, without the caller owning a folder."""
    with bundled_corpus_path(name) as path:
        return load_documents(path)


def chunk_loaded_documents(
    documents: list[LoadedDocument],
    *,
    chunk_size: int,
    chunk_overlap: int,
) -> list[ChunkRecord]:
    """Split loaded documents into ChunkRecords satisfying the metadata contract.

    Boundary-aware: a chunk tries to end on a paragraph/sentence/word break in
    the last 20% of its window rather than mid-word.

    Two properties the caller may rely on:

    **The recorded span is the text.** ``content[char_start:char_end] ==
    text`` for every record. The window is chosen first and the text is then
    stripped of surrounding whitespace, so the *recorded* offsets are the
    stripped span, not the window — a boundary-aligned chunk ends with the
    separator that `strip()` removes, and recording the window there would make
    every offset pair a lie. The stripping is what callers want (no chunk starts
    with a dangling ``"\\n\\n"``); the window still advances on the unstripped
    boundary, so honesty here costs no change in where chunks fall.

    **No chunk is contained in its predecessor.** The last window of a document
    is clamped to the end of the content, which on documents whose length lands
    just past a boundary produced a final chunk of pure overlap — spans
    ``(0,880) (680,1532) (1332,2282) (2082,2282)``, where the fourth is 200
    characters already inside the third. A chunk that says nothing its
    predecessor did not is dropped: it costs an embedding, a row, and a
    duplicate retrieval hit, and buys nothing.
    """
    if chunk_overlap >= chunk_size:
        # Refused, not quietly repaired. Substituting ``chunk_size // 4`` hands
        # the caller a corpus split at parameters they did not ask for and are
        # never told about -- and those parameters are exactly what a stored
        # chunk set claims to be reproducible from.
        raise ValueError(
            f"chunk_overlap ({chunk_overlap}) must be smaller than chunk_size "
            f"({chunk_size}): an overlap at least as wide as the window cannot "
            "advance."
        )

    records: list[ChunkRecord] = []
    global_index = 0
    for doc in documents:
        start = 0
        doc_chunk_index = 0
        content = doc.content or ""
        while start < len(content):
            end = min(start + chunk_size, len(content))
            if end < len(content):
                search_start = start + int(chunk_size * _BOUNDARY_SEARCH_FRACTION)
                best = end
                for sep in _CHUNK_SEPARATORS:
                    pos = content.rfind(sep, search_start, end)
                    if pos > search_start:
                        best = pos + len(sep)
                        break
                end = best
            window = content[start:end]
            text = window.strip()
            if text:
                # Record the span of the STRIPPED text, so
                # content[text_start:text_end] == text holds exactly.
                text_start = start + (len(window) - len(window.lstrip()))
                text_end = text_start + len(text)
                previous = records[-1] if doc_chunk_index else None
                contained_in_previous = (
                    previous is not None
                    and previous.char_start <= text_start
                    and text_end <= previous.char_end
                )
                if not contained_in_previous:
                    chunk_id = f"{doc.source_identity}:chunk-{doc_chunk_index}"
                    records.append(ChunkRecord(
                        chunk_id=chunk_id,
                        doc_id=doc.source_identity,
                        source_file=doc.relative_path,
                        doc_chunk_index=doc_chunk_index,
                        global_index=global_index,
                        text=text,
                        char_start=text_start,
                        char_end=text_end,
                        content_module_id=doc.source_identity,
                        metadata={
                            "source_identity": doc.source_identity,
                            "source_fingerprint": doc.source_fingerprint,
                            "relative_path": doc.relative_path,
                            "filename": doc.filename,
                            "file_type": doc.file_type,
                            "document_filepath": doc.filepath,
                        },
                    ))
                    global_index += 1
                    doc_chunk_index += 1
            next_start = end - chunk_overlap
            if next_start <= start:
                next_start = end
            start = next_start
    return records


# --- the frozen pre-chunked corpus -------------------------------------------
#
# The bundled documents ship with a bundled *chunk set*, and that chunk set is
# the substrate Pelorus's pre-canned extracts and clusters and RAGauge's
# pre-canned dataset and TestSet are all built against. One corpus, one split,
# one set of ids, shared by two products -- which only works if neither has to
# translate the file to read it, so the on-disk shape is exactly the JSONL
# rag-eval already writes: one object per line, these fields, in this order.

#: Top-level JSONL fields, in the order rag-eval writes them. Order is not
#: semantically load-bearing (JSON objects are unordered) but a stable one is
#: what makes the committed file diffable, which is the only way a drifting
#: chunk set is ever noticed in review.
_CHUNK_JSONL_FIELDS = (
    "chunk_id",
    "doc_id",
    "source_file",
    "doc_chunk_index",
    "global_index",
    "char_start",
    "char_end",
    "content_module_id",
    "text",
    "metadata",
)

#: Dropped from the frozen records. ``document_filepath`` is the absolute path
#: the chunker happened to read from -- a temp directory for a zip-imported
#: install, someone's checkout otherwise. Freezing it would commit one
#: machine's filesystem into a file two products read on every other machine,
#: and every consumer already has ``source_file`` and ``relative_path``.
_UNFROZEN_METADATA_KEYS = ("document_filepath",)


def build_bundled_chunks(name: str = DEFAULT_BUNDLED_CORPUS) -> list[ChunkRecord]:
    """Re-derive the pre-chunked corpus from the bundled documents.

    This is the generator behind the committed file. It is public so the
    regeneration script and the test that proves the committed file has not
    drifted from the kit can be the same one line.
    """
    records = chunk_loaded_documents(
        load_bundled_corpus(name),
        chunk_size=BUNDLED_CHUNK_SIZE,
        chunk_overlap=BUNDLED_CHUNK_OVERLAP,
    )
    for record in records:
        for key in _UNFROZEN_METADATA_KEYS:
            record.metadata.pop(key, None)
    return records


def chunks_to_jsonl(records: list[ChunkRecord]) -> str:
    """Render ChunkRecords as the one-object-per-line JSONL both products read.

    ``embedding`` is deliberately not a field: the frozen set is a *chunking*,
    not an index, and a vector is a fact about an embedding model that any
    consumer is free to change.
    """
    lines = []
    for record in records:
        payload = record.to_dict()
        row = {field: payload[field] for field in _CHUNK_JSONL_FIELDS}
        row["metadata"] = dict(row["metadata"] or {})
        lines.append(json.dumps(row, ensure_ascii=False, sort_keys=False))
    return "\n".join(lines) + "\n"


def chunks_from_jsonl(text: str) -> list[ChunkRecord]:
    """Parse the JSONL form back into ChunkRecords.

    Unknown top-level keys are refused rather than ignored: a file carrying a
    field this kit does not model is a file written by something that knows
    something this loader does not, and silently dropping it is how two
    products end up disagreeing about the corpus they both claim to share.
    """
    records: list[ChunkRecord] = []
    for number, line in enumerate(text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        try:
            row = json.loads(line)
        except ValueError as exc:
            raise ValueError(f"chunk record on line {number} is not valid JSON: {exc}") from exc
        unknown = sorted(set(row) - set(_CHUNK_JSONL_FIELDS))
        if unknown:
            raise ValueError(
                f"chunk record on line {number} carries unmodelled field(s) "
                f"{', '.join(unknown)}; this loader would silently drop them."
            )
        records.append(ChunkRecord(
            chunk_id=row["chunk_id"],
            doc_id=row["doc_id"],
            source_file=row["source_file"],
            doc_chunk_index=row["doc_chunk_index"],
            global_index=row.get("global_index"),
            text=row["text"],
            char_start=row.get("char_start"),
            char_end=row.get("char_end"),
            content_module_id=row.get("content_module_id"),
            metadata=dict(row.get("metadata") or {}),
        ))
    return records


@contextmanager
def bundled_chunks_path(name: str = DEFAULT_BUNDLED_CORPUS) -> Iterator[str]:
    """Yield a real file path for the frozen chunk set of a bundled corpus.

    Same ``importlib.resources`` shape, and for the same reason, as
    :func:`bundled_corpus_path`.
    """
    if name not in BUNDLED_CORPORA:
        raise ValueError(
            f"Unknown bundled corpus {name!r}. Available: "
            f"{', '.join(BUNDLED_CORPORA)}."
        )
    resource = (
        importlib.resources.files(__package__)
        .joinpath("corpus")
        .joinpath(f"{name}.chunks.jsonl")
    )
    with importlib.resources.as_file(resource) as path:
        if not Path(path).is_file():
            raise FileNotFoundError(
                f"The frozen chunk set for the {name!r} corpus is not present "
                f"in this installation of rag_connector (looked in {path}). It "
                "is declared in [tool.setuptools.package-data]; a wheel built "
                "without that declaration will not carry it."
            )
        yield str(path)


def load_bundled_chunks(name: str = DEFAULT_BUNDLED_CORPUS) -> list[ChunkRecord]:
    """Load the frozen chunk set: the corpus, already split, ids and all.

    Prefer this over re-chunking the documents whenever anything downstream
    cites a chunk. Re-chunking is a re-derivation, and a re-derivation that
    lands one character differently regrounds every stored citation in both
    products at once; reading the frozen file cannot.
    """
    with bundled_chunks_path(name) as path:
        return chunks_from_jsonl(Path(path).read_text(encoding="utf-8"))
